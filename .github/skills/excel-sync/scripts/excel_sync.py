#!/usr/bin/env python3
"""Two-way sync between user-story-writer markdown artifacts and Excel workbooks.

Markdown (docs/user-story, test-scenario, test-case, test-data) <-> one workbook per type.
Sheet and column names come from ../assets/config.json. See ../SKILL.md for the procedure.

    python excel_sync.py template [--type T] [--out DIR]
    python excel_sync.py status   [--type T] [--id ID]
    python excel_sync.py sync     [--type T] [--id ID] [--prefer md|xlsx]

Requires: pip install openpyxl
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_CONFIG = SKILL_DIR / "assets" / "config.json"
TEMPLATE_DIR = SKILL_DIR / "assets" / "templates"
SYNC_SHEET = "_sync"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
WIDTHS = {
    "story": 50, "goal": 60, "notes": 80, "criterion": 60, "question": 60, "description": 60,
    "action": 50, "expected": 50, "title": 40, "purpose": 40, "preconditions": 50,
    "postconditions": 50, "conventions": 60, "extra": 60, "in_scope": 40, "out_of_scope": 40,
    "related": 40, "linked": 30, "file": 40, "heading": 40, "scenario_title": 40,
}

Q_COLS = [("no", "#"), ("question", "Question"), ("owner", "Owner"), ("status", "Status")]
PRE_COLS = [("no", "#"), ("condition", "Condition"), ("description", "Description")]
CASE_COLS = [("tc_id", "TC ID"), ("title", "Title"), ("type", "Type"), ("technique", "Technique"),
             ("priority", "Priority"), ("ac", "AC")]
STEP_COLS = [("no", "#"), ("action", "Action"), ("expected", "Expected result"), ("test_data", "Test data")]
DATA_INTRO = "Columns after `Class` are the input fields used by the scenario."


# ---------- markdown helpers ----------

def norm_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\r\n?", "\n", str(v)).strip()


def parse_frontmatter(text: str) -> tuple[dict, str]:
    m = re.match(r"^---\r?\n(.*?)\r?\n---\r?\n", text, re.S)
    if not m:
        return {}, text
    meta = {}
    for line in m.group(1).splitlines():
        k, sep, v = line.partition(":")
        if sep:
            meta[k.strip()] = v.strip().strip("\"'")
    return meta, text[m.end():]


def split_h2(body: str) -> tuple[str, list[tuple[str, str]]]:
    h1, sections = "", []
    for line in body.splitlines():
        if line.startswith("# ") and not h1 and not sections:
            h1 = line[2:].strip()
        elif line.startswith("## "):
            sections.append([line[3:].strip(), []])
        elif sections:
            sections[-1][1].append(line)
    return h1, [(h, "\n".join(t).strip()) for h, t in sections]


def split_heading(h1: str, key: str) -> str:
    m = re.search(rf"{re.escape(key)}\s*[-–—]\s*(.*)$", h1)
    return m.group(1).strip() if m else h1


def split_cells(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip().replace("\\|", "|").replace("<br>", "\n") for c in cells]


def table_blocks(text: str) -> list[tuple[int, int]]:
    """Line ranges [start, end) of contiguous table blocks."""
    lines, blocks, start = text.splitlines(), [], None
    for i, line in enumerate(lines + [""]):
        is_row = line.strip().startswith("|")
        if is_row and start is None:
            start = i
        elif not is_row and start is not None:
            blocks.append((start, i))
            start = None
    return blocks


def parse_table(text: str) -> tuple[list[str], list[list[str]]]:
    rows = []
    for line in text.splitlines():
        if not line.strip().startswith("|"):
            continue
        cells = split_cells(line)
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return (rows[0], rows[1:]) if rows else ([], [])


def rows_from_table(text: str, cols: list[tuple[str, str]], ref: str) -> list[dict]:
    keys = [k for k, _ in cols]
    _, rows = parse_table(text)
    return [{"ref": ref, **dict(zip(keys, (r + [""] * len(keys))[:len(keys)]))} for r in rows]


def bullet_items(text: str) -> list[str]:
    return [m.group(1).strip() for m in re.finditer(r"^\s*[-*]\s+(.*)$", text, re.M)]


def parse_bullets(text: str) -> str:
    return "\n".join(bullet_items(text))


def render_table(header: list[str], rows: list[list[str]]) -> str:
    esc = lambda c: str(c).replace("|", "\\|").replace("\n", "<br>")
    lines = ["| " + " | ".join(esc(h) for h in header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
    return "\n".join(lines)


def table(cols: list[tuple[str, str]], rows: list[dict]) -> str:
    return render_table([h for _, h in cols], [[r[k] for k, _ in cols] for r in rows])


def render_bullets(value: str) -> str:
    return "\n".join(f"- {line}" for line in value.splitlines() if line.strip())


def yq(v: str) -> str:
    if v and re.search(r"(^[\s\[\]{}\"'&*!|>%@`#-]|:\s|\s#|\s$)", v):
        return json.dumps(v, ensure_ascii=False)
    return v


def block(title: str, body: str) -> str:
    return f"## {title}\n\n{body}".rstrip() + "\n"


def doc_text(front: list[tuple[str, str]], h1: str, blocks: list[str]) -> str:
    head = "---\n" + "\n".join(f"{k}: {v}" for k, v in front) + "\n---\n"
    return head + f"\n# {h1}\n\n" + "\n".join(blocks).rstrip() + "\n"


# ---------- adapters: markdown <-> canonical doc ----------
# doc = {"key", "main": {...}, "children": {sheet_key: [row, ...]}, "unmapped": [heading, ...]}

AC_RE = re.compile(r"^\*\*(AC-[\w-]+)\*\*\s*(?:\[(.*?)\])?\s*(.*)$", re.S)


def parse_user_story(meta, h1, sections):
    s, key = dict(sections), meta.get("id", "")
    scope = s.get("Scope", "")
    mi = re.search(r"\*\*In scope\*\*(.*?)(?=\*\*Out of scope\*\*|\Z)", scope, re.S)
    mo = re.search(r"\*\*Out of scope\*\*(.*)", scope, re.S)
    criteria = []
    for item in bullet_items(s.get("Acceptance criteria", "")):
        m = AC_RE.match(item)
        criteria.append({"ref": key, "ac": m.group(1), "source": m.group(2) or "", "criterion": m.group(3)}
                        if m else {"ref": key, "ac": "", "source": "", "criterion": item})
    main = {
        "id": key, "title": meta.get("title") or split_heading(h1, key), "status": meta.get("status", ""),
        "requirements": ", ".join(p.strip() for p in meta.get("requirements", "").strip("[]").split(",") if p.strip()),
        "story": s.get("Story", ""), "in_scope": parse_bullets(mi.group(1)) if mi else "",
        "out_of_scope": parse_bullets(mo.group(1)) if mo else "",
        "related": parse_bullets(s.get("Related requirements", "")),
        "linked": parse_bullets(s.get("Linked artifacts", "")), "file": "",
    }
    known = {"Story", "Acceptance criteria", "Scope", "Related requirements", "Linked artifacts", "Open questions"}
    return {"key": key, "main": main, "unmapped": [h for h, _ in sections if h not in known],
            "children": {"criteria": criteria,
                         "questions": rows_from_table(s.get("Open questions", ""), Q_COLS, key)}}


def render_user_story(doc):
    m, c = doc["main"], doc["children"]
    front = [("id", m["id"]), ("type", "user-story"), ("title", yq(m["title"])), ("status", m["status"]),
             ("requirements", f"[{m['requirements']}]")]
    ac = [f"- **{r['ac']}** " + (f"[{r['source']}] " if r["source"] else "") + r["criterion"] for r in c["criteria"]]
    scope = (f"**In scope**\n{render_bullets(m['in_scope'])}\n\n"
             f"**Out of scope**\n{render_bullets(m['out_of_scope'])}")
    return doc_text(front, f"{m['id']} - {m['title']}", [
        block("Story", m["story"]), block("Acceptance criteria", "\n".join(ac)), block("Scope", scope),
        block("Related requirements", render_bullets(m["related"])),
        block("Linked artifacts", render_bullets(m["linked"])),
        block("Open questions", table(Q_COLS, c["questions"])),
    ])


def parse_test_scenario(meta, h1, sections):
    s, key = dict(sections), meta.get("id", "")
    main = {
        "id": key, "story": meta.get("story", ""), "title": meta.get("title", ""),
        "heading": split_heading(h1, key), "status": meta.get("status", ""), "priority": meta.get("priority", ""),
        "goal": s.get("Goal", ""), "test_data": parse_bullets(s.get("Test data", "")),
        "notes": parse_bullets(s.get("Notes", "")), "file": "",
    }
    known = {"Goal", "Covered acceptance criteria", "Preconditions", "Planned test cases", "Test data",
             "Notes", "Open questions"}
    return {"key": key, "main": main, "unmapped": [h for h, _ in sections if h not in known],
            "children": {
                "covered_ac": [{"ref": key, "ac": i} for i in bullet_items(s.get("Covered acceptance criteria", ""))],
                "preconditions": rows_from_table(s.get("Preconditions", ""), PRE_COLS, key),
                "cases": rows_from_table(s.get("Planned test cases", ""), CASE_COLS, key),
                "questions": rows_from_table(s.get("Open questions", ""), Q_COLS, key)}}


def render_test_scenario(doc):
    m, c = doc["main"], doc["children"]
    front = [("id", m["id"]), ("type", "test-scenario"), ("story", m["story"]), ("title", yq(m["title"])),
             ("status", m["status"]), ("priority", m["priority"])]
    return doc_text(front, f"{m['id']} - {m['heading'] or m['title']}", [
        block("Goal", m["goal"]),
        block("Covered acceptance criteria", render_bullets("\n".join(r["ac"] for r in c["covered_ac"]))),
        block("Preconditions", table(PRE_COLS, c["preconditions"])),
        block("Planned test cases", table(CASE_COLS, c["cases"])),
        block("Test data", render_bullets(m["test_data"])),
        block("Notes", render_bullets(m["notes"])),
        block("Open questions", table(Q_COLS, c["questions"])),
    ])


TC_HEAD = re.compile(r"^(TC-[\w-]+)\s*[-–—:]\s*(.*)$")
META_RE = re.compile(r"^\s*-\s+\*\*([^*:]+):\*\*\s*(.*)$", re.M)


def split_blocks(text: str) -> dict[str, str]:
    blocks, cur = {"": []}, ""
    for line in text.splitlines():
        m = re.match(r"^\*\*(Preconditions|Steps|Postconditions)\*\*", line)
        if m:
            cur = m.group(1)
            blocks[cur] = []
        else:
            blocks[cur].append(line)
    return {k: "\n".join(v).strip() for k, v in blocks.items()}


def parse_test_case(meta, h1, sections):
    key = meta.get("scenario", "")
    cases, steps, questions, unmapped = [], [], [], []
    for head, text in sections:
        hm = TC_HEAD.match(head)
        if head == "Open questions":
            questions = rows_from_table(text, Q_COLS, key)
            continue
        if not hm:
            unmapped.append(head)
            continue
        tc_id, title = hm.groups()
        b = split_blocks(text)
        info = {k.strip().lower(): v.strip() for k, v in META_RE.findall(b[""])}
        cases.append({"ref": key, "tc_id": tc_id, "title": title, "ac": info.get("ac", ""),
                      "type": info.get("type", ""), "priority": info.get("priority", ""),
                      "test_data": info.get("test data", ""),
                      "preconditions": parse_bullets(b.get("Preconditions", "")),
                      "postconditions": parse_bullets(b.get("Postconditions", ""))})
        steps += [{**r, "tc_id": tc_id} for r in rows_from_table(b.get("Steps", ""), STEP_COLS, key)]
    main = {"scenario": key, "story": meta.get("story", ""), "scenario_title": split_heading(h1, key),
            "status": meta.get("status", ""), "file": ""}
    return {"key": key, "main": main, "unmapped": unmapped,
            "children": {"cases": cases, "steps": steps, "questions": questions}}


def render_test_case(doc):
    m, c = doc["main"], doc["children"]
    front = [("scenario", m["scenario"]), ("story", m["story"]), ("type", "test-case"), ("status", m["status"])]
    blocks = []
    for r in c["cases"]:
        steps = [s for s in c["steps"] if s["tc_id"] == r["tc_id"]]
        parts = [f"## {r['tc_id']} - {r['title']}\n",
                 f"- **Scenario:** {m['scenario']}\n- **AC:** {r['ac']}\n- **Type:** {r['type']}\n"
                 f"- **Priority:** {r['priority']}\n- **Test data:** {r['test_data']}\n",
                 f"**Preconditions**\n{render_bullets(r['preconditions'])}\n",
                 f"**Steps**\n\n{table(STEP_COLS, steps)}\n"]
        if r["postconditions"]:
            parts.append(f"**Postconditions**\n{render_bullets(r['postconditions'])}\n")
        blocks.append("\n".join(parts))
    if c["questions"]:
        blocks.append(block("Open questions", table(Q_COLS, c["questions"])))
    return doc_text(front, f"Test cases for {m['scenario']} - {m['scenario_title']}", blocks)


def parse_test_data(meta, h1, sections):
    s, key = dict(sections), meta.get("scenario", "")
    text = s.get("Data sets", "")
    lines, tables = text.splitlines(), table_blocks(text)
    header, rows, extra = [], [], ""
    if tables:
        a, b = tables[0]
        header, rows = parse_table("\n".join(lines[a:b]))
        rest = "\n".join(lines[:a] + lines[b:]).strip()
    else:
        rest = text.strip()
    extra = "\n".join(l for l in rest.splitlines() if l.strip() != DATA_INTRO).strip()
    has_used = bool(header) and header[-1].lower() == "used by"
    names = header[3:-1] if has_used else header[3:]
    data = []
    for r in rows:
        r = r + [""] * (len(header) - len(r))
        data.append({"ref": key, "id": r[0], "purpose": r[1], "class": r[2],
                     "used_by": r[-1] if has_used else "", "fields": dict(zip(names, r[3:3 + len(names)]))})
    main = {"scenario": key, "story": meta.get("story", ""), "scenario_title": split_heading(h1, key),
            "status": meta.get("status", ""), "extra": extra,
            "conventions": parse_bullets(s.get("Conventions", "")), "file": ""}
    known = {"Data sets", "Conventions", "Open questions"}
    return {"key": key, "main": main, "unmapped": [h for h, _ in sections if h not in known],
            "children": {"data": data, "questions": rows_from_table(s.get("Open questions", ""), Q_COLS, key)}}


def render_test_data(doc):
    m, c = doc["main"], doc["children"]
    front = [("scenario", m["scenario"]), ("story", m["story"]), ("type", "test-data"), ("status", m["status"])]
    names = []
    for r in c["data"]:
        names += [n for n in r["fields"] if n not in names]
    rows = [[r["id"], r["purpose"], r["class"], *(r["fields"].get(n, "") for n in names), r["used_by"]]
            for r in c["data"]]
    data = f"{DATA_INTRO}\n\n" + render_table(["ID", "Purpose", "Class", *names, "Used by"], rows)
    if m["extra"]:
        data += "\n\n" + m["extra"]
    return doc_text(front, f"Test data for {m['scenario']} - {m['scenario_title']}", [
        block("Data sets", data), block("Conventions", render_bullets(m["conventions"])),
        block("Open questions", table(Q_COLS, c["questions"])),
    ])


ADAPTERS = {
    "user-story": (parse_user_story, render_user_story),
    "test-scenario": (parse_test_scenario, render_test_scenario),
    "test-case": (parse_test_case, render_test_case),
    "test-data": (parse_test_data, render_test_data),
}


def normalize(doc: dict) -> dict:
    doc["main"] = {k: norm_text(v) for k, v in doc["main"].items()}
    for rows in doc["children"].values():
        names = []
        for r in rows:
            for k, v in list(r.items()):
                r[k] = {n: norm_text(x) for n, x in v.items()} if k == "fields" else norm_text(v)
            names += [n for n in r.get("fields", {}) if n not in names]
        used = [n for n in names if any(r["fields"].get(n) for r in rows)]
        for r in rows:
            if "fields" in r:
                r["fields"] = {n: r["fields"].get(n, "") for n in used}
    return doc


def digest(doc: dict) -> str:
    payload = {"main": {k: v for k, v in doc["main"].items() if k != "file"}, "children": doc["children"]}
    return hashlib.sha1(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def parse_md(text: str, tname: str) -> dict:
    meta, body = parse_frontmatter(text)
    h1, sections = split_h2(body)
    return normalize(ADAPTERS[tname][0](meta, h1, sections))


# ---------- workbook layer ----------

def style_sheet(ws, widths: dict[int, int] | None = None) -> None:
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    for i, w in (widths or {}).items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def put_row(ws, r: int, values: list) -> None:
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        if isinstance(v, str) and v.startswith("="):
            cell.data_type = "s"  # keep text starting with "=" from becoming a formula


def build_template(tcfg: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for scfg in tcfg["sheets"].values():
        ws = wb.create_sheet(scfg["name"])
        keys = list(scfg["columns"])
        put_row(ws, 1, list(scfg["columns"].values()) + scfg.get("placeholders", []))
        style_sheet(ws, {i: WIDTHS.get(k, 16) for i, k in enumerate(keys, start=1)})
    wb.create_sheet(SYNC_SHEET).sheet_state = "hidden"
    return wb


def open_workbook(tname: str, tcfg: dict, path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    tpl = TEMPLATE_DIR / f"{tname}.xlsx"
    return load_workbook(tpl) if tpl.exists() else build_template(tcfg)


def read_docs(wb: Workbook, tcfg: dict) -> dict[str, dict]:
    docs: dict[str, dict] = {}
    for skey, scfg in tcfg["sheets"].items():
        if scfg["name"] not in wb.sheetnames:
            sys.exit(f"Sheet '{scfg['name']}' is missing; check assets/config.json or regenerate with 'template'.")
        rows = list(wb[scfg["name"]].iter_rows(values_only=True))
        header = [norm_text(h) for h in rows[0]] if rows else []
        labels = scfg["columns"]
        missing = [lbl for lbl in labels.values() if lbl not in header]
        if missing:
            sys.exit(f"Sheet '{scfg['name']}' lacks columns {missing}; the workbook does not match the config.")
        idx = {k: header.index(lbl) for k, lbl in labels.items()}
        dynamic = scfg.get("dynamic")
        extra = [(i, h) for i, h in enumerate(header) if h and h not in labels.values()] if dynamic else []
        for r in rows[1:]:
            row = {k: norm_text(r[i]) if i < len(r) else "" for k, i in idx.items()}
            if dynamic:
                row["fields"] = {h: norm_text(r[i]) if i < len(r) else "" for i, h in extra}
            if not any(v for k, v in row.items() if k != "fields") and not any(row.get("fields", {}).values()):
                continue
            if skey == "main":
                key = row[tcfg["key_field"]]
                if key:
                    docs[key] = {"key": key, "main": row, "unmapped": [],
                                 "children": {k: [] for k in tcfg["sheets"] if k != "main"}}
            elif row["ref"] in docs:
                docs[row["ref"]]["children"][skey].append(row)
    return {k: normalize(d) for k, d in docs.items()}


def write_docs(wb: Workbook, tcfg: dict, docs: dict[str, dict]) -> None:
    keys = sorted(docs)
    for skey, scfg in tcfg["sheets"].items():
        ws = wb[scfg["name"]]
        ws.delete_rows(1, ws.max_row + 1)
        cols = list(scfg["columns"].items())
        extra: list[str] = []
        if scfg.get("dynamic"):
            for k in keys:
                for r in docs[k]["children"].get(skey, []):
                    extra += [n for n in r.get("fields", {}) if n not in extra]
            extra = extra or scfg.get("placeholders", [])
        put_row(ws, 1, [lbl for _, lbl in cols] + extra)
        n = 1
        for k in keys:
            for r in [docs[k]["main"]] if skey == "main" else docs[k]["children"].get(skey, []):
                n += 1
                put_row(ws, n, [r.get(c, "") for c, _ in cols] + [r.get("fields", {}).get(x, "") for x in extra])
        widths = {i: WIDTHS.get(c, 16) for i, (c, _) in enumerate(cols, start=1)}
        widths.update({len(cols) + i: 18 for i in range(1, len(extra) + 1)})
        style_sheet(ws, widths)


def read_state(wb: Workbook) -> dict[str, str]:
    if SYNC_SHEET not in wb.sheetnames:
        return {}
    return {str(r[0]): str(r[1]) for r in wb[SYNC_SHEET].iter_rows(min_row=2, values_only=True) if r[0]}


def write_state(wb: Workbook, state: dict[str, str]) -> None:
    ws = wb[SYNC_SHEET] if SYNC_SHEET in wb.sheetnames else wb.create_sheet(SYNC_SHEET)
    ws.delete_rows(1, ws.max_row + 1)
    now = datetime.now().isoformat(timespec="seconds")
    put_row(ws, 1, ["Key", "Hash", "Synced at"])
    for i, (k, h) in enumerate(sorted(state.items()), start=2):
        put_row(ws, i, [k, h, now])
    ws.sheet_state = "hidden"


# ---------- sync ----------

def slug(title: str) -> str:
    ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode().lower()
    return "-".join(re.sub(r"[^a-z0-9]+", " ", ascii_title).split()[:5]) or "artifact"


def md_target(root: Path, tcfg: dict, tname: str, doc: dict, existing: dict | None) -> Path:
    folder = (root / tcfg["folder"]).resolve()
    if existing:
        return root / existing["main"]["file"]
    cand = doc["main"].get("file", "")
    if cand:
        p = (root / cand).resolve()
        if p.is_relative_to(folder) and p.suffix == ".md":
            return p
    title = doc["main"].get("title") or doc["main"].get("scenario_title", "")
    return folder / f"{doc['key']}-{slug(title)}{tcfg['glob'].lstrip('*')}"


def load_md(root: Path, tcfg: dict, tname: str) -> dict[str, dict]:
    docs = {}
    for path in sorted((root / tcfg["folder"]).glob(tcfg["glob"])):
        doc = parse_md(path.read_text(encoding="utf-8"), tname)
        if not doc["key"]:
            print(f"WARN  {path.name}: no key in frontmatter, skipped")
            continue
        doc["main"]["file"] = path.relative_to(root).as_posix()
        docs[doc["key"]] = doc
    return docs


def plan(m, x, base, prefer) -> str:
    hm, hx = (digest(m) if m else None), (digest(x) if x else None)
    if m and x:
        if hm == hx:
            return "in-sync"
        if base == hx:
            return "push"
        if base == hm:
            return "pull"
        return {"md": "push", "xlsx": "pull"}.get(prefer, "conflict")
    if m:
        return "push" if base is None or prefer == "md" else "xlsx-removed"
    return "pull" if base is None or prefer == "xlsx" else "md-removed"


def run_type(root: Path, cfg: dict, tname: str, args, apply: bool) -> int:
    tcfg = cfg["types"][tname]
    path = root / cfg["workbook_dir"] / tcfg["workbook"]
    wb = open_workbook(tname, tcfg, path)
    md, xd, state = load_md(root, tcfg, tname), read_docs(wb, tcfg), read_state(wb)
    new_state, dirty, problems = dict(state), False, 0
    for key in sorted(set(md) | set(xd)):
        if args.id and key != args.id:
            continue
        m, x = md.get(key), xd.get(key)
        action = plan(m, x, state.get(key), args.prefer)
        note = ""
        if action == "pull" and m and m["unmapped"] and args.prefer != "xlsx":
            action, note = "blocked", f"md has sections not mapped to xlsx {m['unmapped']}; pull would drop them"
        if action in ("conflict", "blocked", "xlsx-removed", "md-removed"):
            problems += 1
        print(f"{tname:14} {key:14} {action:13} {note}".rstrip())
        if not apply:
            continue
        if action == "push":
            xd[key], dirty = m, True
            new_state[key] = digest(m)
        elif action == "pull":
            target = md_target(root, tcfg, tname, x, m)
            text = ADAPTERS[tname][1](x)
            target.parent.mkdir(parents=True, exist_ok=True)
            crlf = target.exists() and b"\r\n" in target.read_bytes()
            with open(target, "w", encoding="utf-8", newline="\r\n" if crlf else "\n") as f:
                f.write(text)
            if digest(parse_md(text, tname)) != digest(x):
                print(f"WARN  {target.name}: markdown does not round-trip to the xlsx content; review it")
            x["main"]["file"], dirty = target.relative_to(root).as_posix(), True
            new_state[key] = digest(x)
        elif action == "in-sync":
            new_state[key] = digest(m)
    if apply and (dirty or new_state != state):
        write_docs(wb, tcfg, xd)
        write_state(wb, new_state)
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            wb.save(path)
        except PermissionError:
            sys.exit(f"Cannot write {path}; close it in Excel and run again.")
    return problems


def cmd_template(root: Path, cfg: dict, args) -> int:
    out = Path(args.out) if args.out else TEMPLATE_DIR
    out.mkdir(parents=True, exist_ok=True)
    for tname in args.types:
        build_template(cfg["types"][tname]).save(out / f"{tname}.xlsx")
        print(f"Saved {out / (tname + '.xlsx')}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--root", type=Path, default=SKILL_DIR.parents[2], help="workspace root")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    sub = parser.add_subparsers(dest="command", required=True)
    for name in ("template", "status", "sync"):
        p = sub.add_parser(name)
        p.add_argument("--type", action="append", dest="types", help="artifact type (repeatable); default: all")
        if name == "template":
            p.add_argument("--out", help="output directory (default: assets/templates)")
        else:
            p.add_argument("--id", help="only this artifact key (e.g. US-KOS-01, TS-KOS-01)")
        if name == "sync":
            p.add_argument("--prefer", choices=["md", "xlsx"], help="resolve conflicts and removals in favour of one side")
    args = parser.parse_args()
    args.prefer = getattr(args, "prefer", None)
    args.id = getattr(args, "id", None)

    cfg = json.loads(args.config.read_text(encoding="utf-8"))
    unknown = [t for t in args.types or [] if t not in cfg["types"]]
    if unknown:
        sys.exit(f"Unknown type(s) {unknown}; available: {list(cfg['types'])}")
    args.types = args.types or list(cfg["types"])
    if args.command == "template":
        return cmd_template(args.root, cfg, args)
    problems = sum(run_type(args.root, cfg, t, args, apply=args.command == "sync") for t in args.types)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())

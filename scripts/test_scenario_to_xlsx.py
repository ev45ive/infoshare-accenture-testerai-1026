#!/usr/bin/env python3
"""Convert test scenario markdown files (docs/test-scenario) to one Excel workbook.

Usage:
    python scripts/test_scenario_to_xlsx.py                      # all of docs/test-scenario
    python scripts/test_scenario_to_xlsx.py FILE_OR_DIR ... -o out.xlsx

Requires: pip install openpyxl
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = ROOT / "docs" / "test-scenario"
DEFAULT_OUTPUT = ROOT / "docs" / "test-scenario" / "test-scenarios.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    match = re.match(r"^---\r?\n(.*?)\r?\n---\r?\n", text, re.DOTALL)
    if not match:
        return {}, text
    meta = {}
    for line in match.group(1).splitlines():
        key, sep, value = line.partition(":")
        if sep:
            meta[key.strip()] = value.strip().strip("\"'")
    return meta, text[match.end():]


def split_sections(body: str) -> tuple[str, dict[str, str]]:
    """Return the H1 title and a map of H2 heading -> section text."""
    title = ""
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in body.splitlines():
        if line.startswith("# ") and not title:
            title = line[2:].strip()
        elif line.startswith("## "):
            current = line[3:].strip()
            sections[current] = []
        elif current is not None:
            sections[current].append(line)
    return title, {k: "\n".join(v).strip() for k, v in sections.items()}


def parse_table(text: str) -> tuple[list[str], list[list[str]]]:
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", line.strip("|"))]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append([c.replace("\\|", "|") for c in cells])
    return (rows[0], rows[1:]) if rows else ([], [])


def parse_bullets(text: str) -> list[str]:
    return [m.group(1).strip() for m in re.finditer(r"^\s*[-*]\s+(.*)$", text, re.MULTILINE)]


def strip_md(value: str) -> str:
    return re.sub(r"[`*]", "", value)


class Scenario:
    def __init__(self, path: Path):
        self.path = path
        meta, body = parse_frontmatter(path.read_text(encoding="utf-8"))
        self.meta = meta
        self.title, self.sections = split_sections(body)
        self.id = meta.get("id", path.stem)

    def table(self, name: str) -> tuple[list[str], list[list[str]]]:
        return parse_table(self.sections.get(name, ""))

    def bullets(self, name: str) -> list[str]:
        return parse_bullets(self.sections.get(name, ""))


def add_sheet(wb: Workbook, name: str, header: list[str], rows: list[list[str]], widths: list[int]) -> None:
    ws = wb.create_sheet(name)
    ws.append(header)
    for row in rows:
        ws.append([strip_md(str(c)) for c in row])
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(scenarios: list[Scenario]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    overview = [
        [s.id, s.meta.get("story", ""), s.meta.get("title", s.title), s.meta.get("status", ""),
         s.meta.get("priority", ""), s.sections.get("Goal", ""), "\n".join(s.bullets("Notes")), s.path.name]
        for s in scenarios
    ]
    add_sheet(wb, "Scenarios",
              ["Scenario ID", "Story", "Title", "Status", "Priority", "Goal", "Notes", "Source file"],
              overview, [14, 14, 40, 12, 10, 60, 80, 40])

    cases, ac_rows, pre_rows, td_rows, oq_rows = [], [], [], [], []
    case_header: list[str] = []
    for s in scenarios:
        header, rows = s.table("Planned test cases")
        case_header = case_header or header
        cases += [[s.id] + r for r in rows]
        ac_rows += [[s.id, b] for b in s.bullets("Covered acceptance criteria")]
        _, rows = s.table("Preconditions")
        pre_rows += [[s.id] + r for r in rows]
        td_rows += [[s.id, b] for b in s.bullets("Test data")]
        _, rows = s.table("Open questions")
        oq_rows += [[s.id] + r for r in rows]

    add_sheet(wb, "Test Cases", ["Scenario ID"] + (case_header or ["TC ID", "Title", "Type", "Technique", "Priority", "AC"]),
              cases, [14, 16, 60, 12, 22, 10, 10])
    add_sheet(wb, "Covered AC", ["Scenario ID", "Acceptance criterion"], ac_rows, [14, 100])
    add_sheet(wb, "Preconditions", ["Scenario ID", "#", "Condition", "Description"], pre_rows, [14, 5, 16, 100])
    add_sheet(wb, "Test Data", ["Scenario ID", "Test data"], td_rows, [14, 100])
    add_sheet(wb, "Open Questions", ["Scenario ID", "#", "Question", "Owner", "Status"], oq_rows, [14, 5, 100, 18, 10])
    return wb


def collect_files(inputs: list[Path]) -> list[Path]:
    files: list[Path] = []
    for p in inputs:
        files += sorted(p.rglob("*.test-scenario.md")) if p.is_dir() else [p]
    return files


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("inputs", nargs="*", type=Path, default=[DEFAULT_INPUT], help="scenario files or directories")
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT, help="output .xlsx path")
    args = parser.parse_args()

    files = collect_files(args.inputs)
    if not files:
        print("No test scenario files found.", file=sys.stderr)
        return 1

    wb = build_workbook([Scenario(f) for f in files])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(f"Saved {args.output} ({len(files)} scenario file(s))")
    return 0


if __name__ == "__main__":
    sys.exit(main())
